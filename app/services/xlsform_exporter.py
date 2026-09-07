from pathlib import Path

from openpyxl import load_workbook
from werkzeug.utils import secure_filename

from app.models import TranslationReview
from app.services.review_workflow import choice_items_by_list_name, displayed_survey_items_query
from app.services.xlsform_parser import (
    ENGLISH_LABEL_HEADER,
    TRANSLATABLE_FIELDS,
    find_localized_header,
    localized_header,
    reconstruct_translation_cell,
)


class XLSFormExportError(Exception):
    """Raised when a reviewed XLSForm cannot be generated safely."""


def export_reviewed_xlsform(assignment, upload_folder, export_folder):
    source_path = (Path(upload_folder).expanduser().resolve() / assignment.xlsform.stored_filename).resolve()
    if not source_path.is_file():
        raise XLSFormExportError("The original uploaded XLSForm could not be found.")

    export_dir = Path(export_folder).expanduser().resolve()
    export_dir.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(source_path)
    allowed_keys = _exportable_review_keys(assignment)
    reviews = TranslationReview.query.filter_by(
        xlsform_id=assignment.xlsform_id,
        language_id=assignment.language_id,
        edited_by=assignment.user_id,
    ).all()
    reviews_by_sheet_row = {
        (review.sheet_name, review.row_number, review.field_name): review
        for review in reviews
        if (review.sheet_name, review.row_number, review.field_name) in allowed_keys
    }

    for sheet_name in ("survey", "choices"):
        sheet = _sheet_by_name(workbook, sheet_name)
        if sheet is None:
            raise XLSFormExportError(f"The original XLSForm is missing the {sheet_name} sheet.")
        field_names = TRANSLATABLE_FIELDS if sheet_name == "survey" else ("label",)
        headers = [cell.value for cell in sheet[1]]
        if sheet_name == "survey":
            label_header = find_localized_header(headers, "label", assignment.language.excel_header)
            if _column_for_header(sheet, label_header) is None:
                raise XLSFormExportError(
                    f"The target language column {assignment.language.excel_header!r} is missing from {sheet_name}."
                )
        for field_name in field_names:
            field_reviews = {
                row_number: review
                for (review_sheet, row_number, review_field), review in reviews_by_sheet_row.items()
                if review_sheet == sheet_name and review_field == field_name
            }
            if not field_reviews:
                continue
            target_header = find_localized_header(headers, field_name, assignment.language.excel_header)
            target_column = _column_for_header(sheet, target_header)
            if target_column is None:
                if field_name == "label":
                    raise XLSFormExportError(
                        f"The target language column {target_header!r} is missing from {sheet_name}."
                    )
                target_column = sheet.max_column + 1
                target_header = localized_header(field_name, assignment.language.excel_header)
                sheet.cell(row=1, column=target_column).value = target_header
                headers.append(target_header)
            english_field_header = find_localized_header(headers, field_name, ENGLISH_LABEL_HEADER)
            english_column = _column_for_header(sheet, english_field_header)
            for row_number, review in field_reviews.items():
                english_reference = sheet.cell(row=row_number, column=english_column).value if english_column else None
                sheet.cell(row=row_number, column=target_column).value = reconstruct_translation_cell(
                    english_reference,
                    review.original_cell_value,
                    review.edited_translation,
                )

    export_filename = _export_filename(assignment)
    export_path = export_dir / export_filename
    workbook.save(export_path)
    if not export_path.is_file():
        raise XLSFormExportError("The reviewed XLSForm was not created.")
    return export_path


def _sheet_by_name(workbook, sheet_name):
    for actual_name in workbook.sheetnames:
        if actual_name.casefold() == sheet_name.casefold():
            return workbook[actual_name]
    return None


def _column_for_header(sheet, header):
    for cell in sheet[1]:
        if cell.value == header:
            return cell.column
    return None


def _exportable_review_keys(assignment):
    keys = set()
    choices_by_list = choice_items_by_list_name(assignment.xlsform)
    for item in displayed_survey_items_query(assignment.xlsform).all():
        for field_name in TRANSLATABLE_FIELDS:
            keys.add(("survey", item.row_number, field_name))
        for choice in choices_by_list.get(item.list_name, []):
            keys.add(("choices", choice.row_number, "label"))
    return keys


def _export_filename(assignment):
    original_stem = Path(assignment.xlsform.original_filename).stem
    code = assignment.language.language_code or assignment.language.display_name
    filename = f"{original_stem}_{code}_{assignment.reviewer.username}_reviewed.xlsx"
    return secure_filename(filename)
