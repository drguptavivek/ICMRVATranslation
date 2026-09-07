import re
from dataclasses import dataclass
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


REQUIRED_SHEETS = ("survey", "choices", "settings")
TRANSLATABLE_FIELDS = (
    "label",
    "hint",
    "constraint_message",
    "required_message",
    "guidance_hint",
)
LANGUAGE_HEADER_PREFIX = "label::"
ENGLISH_LABEL_HEADER = "label::English (en)"
LANGUAGE_RE = re.compile(
    rf"^(?P<field>{'|'.join(TRANSLATABLE_FIELDS)})::(?P<name>.+?)(?:\s*\((?P<code>[^()]*)\)\s*)?$",
    re.IGNORECASE,
)


class XLSFormValidationError(Exception):
    pass


@dataclass(frozen=True)
class DetectedLanguage:
    display_name: str
    language_code: str | None
    excel_header: str


@dataclass(frozen=True)
class XLSFormParseResult:
    form_title: str | None
    form_id: str | None
    version: str | None
    survey_row_count: int
    choices_row_count: int
    languages: list[DetectedLanguage]
    survey_items: list["ParsedSurveyItem"]
    choice_items: list["ParsedChoiceItem"]


@dataclass(frozen=True)
class ParsedSurveyItem:
    row_number: int
    type: str | None
    name: str | None
    list_name: str | None
    english_label: str | None
    relevant: str | None
    appearance: str | None
    is_group: bool
    raw_row_data: dict


@dataclass(frozen=True)
class ParsedChoiceItem:
    row_number: int
    list_name: str | None
    name: str | None
    english_label: str | None
    raw_row_data: dict


@dataclass(frozen=True)
class ParsedTranslation:
    original_cell_value: str | None
    extracted_translation: str | None


def parse_xlsform(file_obj):
    try:
        workbook = load_workbook(file_obj, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError, KeyError, ValueError) as exc:
        raise XLSFormValidationError(
            "The uploaded file could not be read as a valid XLSForm workbook."
        ) from exc

    if not workbook.sheetnames:
        raise XLSFormValidationError("The uploaded workbook is empty.")

    normalized_sheets = {sheet_name.lower(): sheet_name for sheet_name in workbook.sheetnames}
    missing_sheets = [name for name in REQUIRED_SHEETS if name not in normalized_sheets]
    if missing_sheets:
        missing = ", ".join(missing_sheets)
        raise XLSFormValidationError(f"The uploaded XLSForm is missing required sheet(s): {missing}.")

    survey = workbook[normalized_sheets["survey"]]
    choices = workbook[normalized_sheets["choices"]]
    settings = workbook[normalized_sheets["settings"]]

    survey_headers = _read_headers(survey, "survey")
    choices_headers = _read_headers(choices, "choices")
    settings_metadata = _read_settings_metadata(settings)

    languages = _detect_languages(survey_headers + choices_headers)
    survey_items = _read_survey_items(survey, survey_headers)
    choice_items = _read_choice_items(choices, choices_headers)

    return XLSFormParseResult(
        form_title=settings_metadata.get("form_title"),
        form_id=settings_metadata.get("form_id"),
        version=settings_metadata.get("version"),
        survey_row_count=_count_data_rows(survey),
        choices_row_count=_count_data_rows(choices),
        languages=languages,
        survey_items=survey_items,
        choice_items=choice_items,
    )


def parse_translation_cell(english_reference, original_cell_value):
    original = _clean_cell(original_cell_value)
    english = _clean_cell(english_reference)
    if not original:
        return ParsedTranslation(original_cell_value=original, extracted_translation=None)
    if english and original == english:
        return ParsedTranslation(original_cell_value=original, extracted_translation="")
    if english and original.startswith(english):
        remaining = original[len(english):]
        if remaining.startswith("\r\n"):
            return ParsedTranslation(
                original_cell_value=original,
                extracted_translation=remaining[2:].strip(),
            )
        if remaining.startswith("\n") or remaining.startswith("\r"):
            return ParsedTranslation(
                original_cell_value=original,
                extracted_translation=remaining[1:].strip(),
            )
    return ParsedTranslation(original_cell_value=original, extracted_translation=original)


def reconstruct_translation_cell(english_reference, original_cell_value, edited_translation):
    original = _clean_cell(original_cell_value)
    english = _clean_cell(english_reference)
    edited = (edited_translation or "").strip()
    if not english:
        return edited
    if original and _cell_starts_with_english_reference(english, original):
        return f"{english}\n{edited}" if edited else english
    return edited


def _read_headers(sheet, sheet_name):
    rows = sheet.iter_rows(min_row=1, max_row=1, values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        raise XLSFormValidationError(f"The {sheet_name} sheet is empty.") from exc

    headers = [_clean_cell(value) for value in header_row]
    headers = [header for header in headers if header]
    if not headers:
        raise XLSFormValidationError(f"The {sheet_name} sheet does not contain a header row.")
    return headers


def _read_settings_metadata(sheet):
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [_clean_cell(value) for value in next(rows)]
    except StopIteration:
        return {}

    try:
        values = next(rows)
    except StopIteration:
        return {}

    metadata = {}
    for index, header in enumerate(headers):
        if header in {"form_title", "form_id", "version"}:
            metadata[header] = _clean_cell(values[index] if index < len(values) else None)
    return metadata


def _read_survey_items(sheet, headers):
    items = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        raw_row_data = _row_to_dict(headers, row)
        if not any(raw_row_data.values()):
            continue
        type_value = raw_row_data.get("type")
        items.append(
            ParsedSurveyItem(
                row_number=row_number,
                type=type_value,
                name=raw_row_data.get("name"),
                list_name=_survey_list_name(type_value),
                english_label=raw_row_data.get(ENGLISH_LABEL_HEADER),
                relevant=raw_row_data.get("relevant"),
                appearance=raw_row_data.get("appearance"),
                is_group=type_value in {"begin_group", "begin repeat"},
                raw_row_data=raw_row_data,
            )
        )
    return items


def _read_choice_items(sheet, headers):
    items = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        raw_row_data = _row_to_dict(headers, row)
        if not any(raw_row_data.values()):
            continue
        items.append(
            ParsedChoiceItem(
                row_number=row_number,
                list_name=raw_row_data.get("list_name"),
                name=raw_row_data.get("name"),
                english_label=raw_row_data.get(ENGLISH_LABEL_HEADER),
                raw_row_data=raw_row_data,
            )
        )
    return items


def _count_data_rows(sheet):
    count = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if any(_clean_cell(value) for value in row):
            count += 1
    return count


def _detect_languages(headers):
    detected = []
    seen = set()
    for header in headers:
        match = LANGUAGE_RE.match(header)
        if not match:
            continue
        language = _parse_language_header(header)
        key = ((language.language_code or "").casefold(), language.display_name.casefold())
        if key in seen:
            continue
        seen.add(key)
        detected.append(language)
    return detected


def _row_to_dict(headers, row):
    raw_row_data = {}
    for index, header in enumerate(headers):
        raw_row_data[header] = _clean_cell(row[index] if index < len(row) else None)
    return raw_row_data


def _survey_list_name(type_value):
    if not type_value:
        return None
    parts = type_value.split()
    if len(parts) >= 2 and parts[0] in {"select_one", "select_multiple"}:
        return parts[1]
    return None


def _parse_language_header(header):
    match = LANGUAGE_RE.match(header)
    if not match:
        display_name = header.split("::", 1)[-1].strip()
        return DetectedLanguage(
            display_name=display_name,
            language_code=None,
            excel_header=f"{LANGUAGE_HEADER_PREFIX}{display_name}",
        )

    display_name = match.group("name").strip()
    language_code = match.group("code")
    if language_code:
        language_code = language_code.strip() or None
    return DetectedLanguage(
        display_name=display_name,
        language_code=language_code,
        excel_header=(
            header
            if match.group("field").casefold() == "label"
            else f"{LANGUAGE_HEADER_PREFIX}{header.split('::', 1)[-1]}"
        ),
    )


def localized_header(field_name, label_header):
    suffix = label_header.split("::", 1)[-1]
    return f"{field_name}::{suffix}"


def english_header(field_name):
    return localized_header(field_name, ENGLISH_LABEL_HEADER)


def find_localized_header(headers, field_name, label_header):
    target = LANGUAGE_RE.match(label_header)
    if target:
        target_name = target.group("name").strip().casefold()
        target_code = (target.group("code") or "").strip().casefold()
        for header in headers:
            if not isinstance(header, str):
                continue
            match = LANGUAGE_RE.match(header)
            if not match or match.group("field").casefold() != field_name.casefold():
                continue
            name = match.group("name").strip().casefold()
            code = (match.group("code") or "").strip().casefold()
            if target_code and code == target_code:
                return header
            if not target_code and name == target_name:
                return header
    return localized_header(field_name, label_header)


def _cell_starts_with_english_reference(english, original):
    if original == english:
        return True
    remaining = original[len(english):] if original.startswith(english) else None
    return remaining is not None and (
        remaining.startswith("\r\n")
        or remaining.startswith("\n")
        or remaining.startswith("\r")
    )


def _clean_cell(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None
