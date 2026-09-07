import io
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook

from app import create_app
from app.extensions import db
from app.models import ChoiceItem, Language, SurveyItem, TranslationReview, User, XLSForm
from app.services.xlsform_parser import XLSFormValidationError, parse_xlsform


class TestConfig:
    TESTING = True
    SECRET_KEY = "test-secret"
    SQLALCHEMY_DATABASE_URI = "sqlite:///:memory:"
    SQLALCHEMY_TRACK_MODIFICATIONS = False
    WTF_CSRF_ENABLED = False
    UPLOAD_FOLDER = "instance/test_uploads"
    EXPORT_FOLDER = "instance/test_exports"


class XLSFormUploadTestCase(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.app = create_app(TestConfig)
        self.app.config["UPLOAD_FOLDER"] = str(Path(self.temp_dir.name) / "uploads")
        self.client = self.app.test_client()
        self.app_context = self.app.app_context()
        self.app_context.push()
        db.create_all()
        self.admin = self.create_user("admin", "admin@example.com", User.ROLE_ADMIN)
        self.reviewer = self.create_user("reviewer", "reviewer@example.com", User.ROLE_REVIEWER)

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.app_context.pop()
        self.temp_dir.cleanup()

    def create_user(self, username, email, role, password="password123"):
        user = User(username=username, email=email, role=role, is_active=True)
        user.set_password(password)
        db.session.add(user)
        db.session.commit()
        return user

    def login(self, username):
        return self.client.post(
            "/login",
            data={"username_or_email": username, "password": "password123"},
        )

    def test_valid_source_xlsform_upload(self):
        self.login("admin")
        source_path = Path("source_files/va_who_2022_All_Language_Training-Module.xlsx")
        with source_path.open("rb") as source_file:
            response = self.client.post(
                "/admin/xlsforms",
                data={"xlsform_file": (source_file, source_path.name)},
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Import Summary", response.data)
        self.assertEqual(XLSForm.query.count(), 1)
        self.assertGreater(Language.query.count(), 0)
        stored_filename = XLSForm.query.first().stored_filename
        self.assertTrue((Path(self.app.config["UPLOAD_FOLDER"]) / stored_filename).exists())

    def test_missing_required_sheet_is_rejected(self):
        self.login("admin")
        workbook_bytes = make_workbook(include_choices=False)
        response = self.client.post(
            "/admin/xlsforms",
            data={"xlsform_file": (workbook_bytes, "missing_choices.xlsx")},
            content_type="multipart/form-data",
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"missing required sheet(s): choices", response.data)
        self.assertEqual(XLSForm.query.count(), 0)

    def test_invalid_file_type_is_rejected(self):
        self.login("admin")
        response = self.client.post(
            "/admin/xlsforms",
            data={"xlsform_file": (io.BytesIO(b"not an xlsx"), "form.txt")},
            content_type="multipart/form-data",
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Upload a valid .xlsx file.", response.data)
        self.assertEqual(XLSForm.query.count(), 0)

    def test_language_detection(self):
        result = parse_xlsform(make_workbook())

        self.assertEqual(result.form_title, "Training Form")
        self.assertEqual(result.form_id, "training_form")
        self.assertEqual(result.version, "2026-09-01")
        self.assertEqual(result.survey_row_count, 1)
        self.assertEqual(result.choices_row_count, 1)
        self.assertEqual(
            [(language.display_name, language.language_code, language.excel_header) for language in result.languages],
            [
                ("English", "en", "label::English (en)"),
                ("Punjabi", "pu", "label::Punjabi(pu)"),
                ("Hindi", "hi", "label::Hindi (hi)"),
            ],
        )

    def test_duplicate_language_prevention(self):
        self.login("admin")
        for _ in range(2):
            self.client.post(
                "/admin/xlsforms",
                data={"xlsform_file": (make_workbook(), "form.xlsx")},
                content_type="multipart/form-data",
            )

        self.assertEqual(XLSForm.query.count(), 2)
        self.assertEqual(Language.query.count(), 3)

    def test_admin_authorization_required(self):
        response = self.client.get("/admin/xlsforms")
        self.assertEqual(response.status_code, 302)
        self.assertIn("/login", response.headers["Location"])

        self.login("reviewer")
        response = self.client.get("/admin/xlsforms")
        self.assertEqual(response.status_code, 403)

    def test_existing_xlsform_content_import_is_idempotent(self):
        self.login("admin")
        parse_result = parse_xlsform(make_workbook())
        xlsform = XLSForm(
            original_filename="existing.xlsx",
            stored_filename="existing.xlsx",
            uploaded_by=self.admin.id,
            form_title=parse_result.form_title,
            form_id=parse_result.form_id,
            version=parse_result.version,
            survey_row_count=parse_result.survey_row_count,
            choices_row_count=parse_result.choices_row_count,
            detected_language_count=len(parse_result.languages),
        )
        upload_dir = Path(self.app.config["UPLOAD_FOLDER"])
        upload_dir.mkdir(parents=True, exist_ok=True)
        workbook_bytes = make_workbook()
        (upload_dir / xlsform.stored_filename).write_bytes(workbook_bytes.read())
        db.session.add(xlsform)
        db.session.commit()

        for _ in range(2):
            response = self.client.post(
                f"/admin/xlsforms/{xlsform.id}/import",
                follow_redirects=True,
            )
            self.assertEqual(response.status_code, 200)
            self.assertIn(b"Imported", response.data)

        self.assertEqual(SurveyItem.query.filter_by(xlsform_id=xlsform.id).count(), 1)
        self.assertEqual(ChoiceItem.query.filter_by(xlsform_id=xlsform.id).count(), 1)

    def test_reimport_preserves_saved_translation_edits(self):
        self.login("admin")
        parse_result = parse_xlsform(make_workbook())
        xlsform = XLSForm(
            original_filename="existing.xlsx",
            stored_filename="existing.xlsx",
            uploaded_by=self.admin.id,
            form_title=parse_result.form_title,
            form_id=parse_result.form_id,
            version=parse_result.version,
            survey_row_count=parse_result.survey_row_count,
            choices_row_count=parse_result.choices_row_count,
            detected_language_count=len(parse_result.languages),
        )
        upload_dir = Path(self.app.config["UPLOAD_FOLDER"])
        upload_dir.mkdir(parents=True, exist_ok=True)
        workbook_bytes = make_workbook()
        (upload_dir / xlsform.stored_filename).write_bytes(workbook_bytes.read())
        db.session.add(xlsform)
        db.session.commit()

        self.client.post(f"/admin/xlsforms/{xlsform.id}/import")
        hindi = Language.query.filter_by(excel_header="label::Hindi (hi)").first()
        review = TranslationReview.query.filter_by(
            xlsform_id=xlsform.id,
            sheet_name="choices",
            row_number=2,
            language_id=hindi.id,
        ).first()
        review.edited_translation = "Saved edit"
        review.edited_by = self.admin.id
        db.session.commit()

        self.client.post(f"/admin/xlsforms/{xlsform.id}/import")
        review = TranslationReview.query.filter_by(
            xlsform_id=xlsform.id,
            sheet_name="choices",
            row_number=2,
            language_id=hindi.id,
        ).first()
        self.assertEqual(review.edited_translation, "Saved edit")

    def test_new_upload_matches_reviews_by_identity_and_updates_untouched_rows(self):
        self.login("admin")
        self.client.post(
            "/admin/xlsforms",
            data={"xlsform_file": (make_continuity_workbook(), "first.xlsx")},
            content_type="multipart/form-data",
        )
        first = XLSForm.query.order_by(XLSForm.id).first()
        punjabi = Language.query.filter_by(excel_header="label::Punjabi(pu)").first()
        hindi = Language.query.filter_by(excel_header="label::Hindi (hi)").first()

        survey = survey_review(first, "q1", punjabi.id)
        survey.edited_translation = "Saved q1"
        survey.edited_by = self.admin.id
        survey.edited_at = datetime.now(timezone.utc)
        survey.reviewed_at = datetime.now(timezone.utc)
        survey.status = TranslationReview.STATUS_REVIEWED
        choice = choice_review(first, "list_a", "shared", hindi.id)
        choice.edited_translation = "Saved list A"
        choice.edited_by = self.admin.id
        choice.edited_at = datetime.now(timezone.utc)
        db.session.commit()

        self.client.post(
            "/admin/xlsforms",
            data={"xlsform_file": (make_continuity_workbook(second=True), "second.xlsx")},
            content_type="multipart/form-data",
        )
        second = XLSForm.query.order_by(XLSForm.id.desc()).first()

        survey = survey_review(second, "q1", punjabi.id)
        self.assertEqual(survey.edited_translation, "Saved q1")
        self.assertEqual(survey.status, TranslationReview.STATUS_REVIEWED)
        self.assertEqual(survey.original_cell_value, "Naam updated")
        self.assertEqual(survey_review(second, "q2", punjabi.id).edited_translation, "Updated age")
        self.assertIsNotNone(survey_review(second, "q3", punjabi.id))
        self.assertEqual(
            choice_review(second, "list_a", "shared", hindi.id).edited_translation,
            "Saved list A",
        )
        self.assertEqual(
            choice_review(second, "list_b", "shared", hindi.id).edited_translation,
            "Updated list B",
        )

    def test_reimport_resets_approval_when_english_source_changes(self):
        self.login("admin")
        self.client.post(
            "/admin/xlsforms",
            data={"xlsform_file": (make_continuity_workbook(), "source.xlsx")},
            content_type="multipart/form-data",
        )
        xlsform = XLSForm.query.first()
        punjabi = Language.query.filter_by(excel_header="label::Punjabi(pu)").first()
        review = survey_review(xlsform, "q1", punjabi.id)
        review.edited_translation = "Saved q1"
        review.edited_by = self.admin.id
        review.edited_at = datetime.now(timezone.utc)
        review.reviewed_at = datetime.now(timezone.utc)
        review.status = TranslationReview.STATUS_REVIEWED
        db.session.commit()

        upload_path = Path(self.app.config["UPLOAD_FOLDER"]) / xlsform.stored_filename
        upload_path.write_bytes(make_continuity_workbook(english_changed=True).read())
        response = self.client.post(f"/admin/xlsforms/{xlsform.id}/import")
        self.assertEqual(response.status_code, 302)

        review = survey_review(xlsform, "q1", punjabi.id)
        self.assertEqual(review.edited_translation, "Saved q1")
        self.assertEqual(review.status, TranslationReview.STATUS_PENDING)
        self.assertIsNone(review.reviewed_at)
        self.assertEqual(review.original_cell_value, "Naam changed")

def survey_review(xlsform, name, language_id):
    item = SurveyItem.query.filter_by(xlsform_id=xlsform.id, name=name).one()
    return TranslationReview.query.filter_by(
        xlsform_id=xlsform.id,
        sheet_name="survey",
        row_number=item.row_number,
        field_name="label",
        language_id=language_id,
    ).one()


def choice_review(xlsform, list_name, name, language_id):
    item = ChoiceItem.query.filter_by(
        xlsform_id=xlsform.id,
        list_name=list_name,
        name=name,
    ).one()
    return TranslationReview.query.filter_by(
        xlsform_id=xlsform.id,
        sheet_name="choices",
        row_number=item.row_number,
        field_name="label",
        language_id=language_id,
    ).one()


def make_workbook(include_choices=True):
    workbook = Workbook()
    survey = workbook.active
    survey.title = "survey"
    survey.append(["type", "name", "label::English (en)", "label::Punjabi(pu)"])
    survey.append(["text", "name", "Name", "Naam"])

    if include_choices:
        choices = workbook.create_sheet("choices")
        choices.append(["list_name", "name", "label::Hindi (hi)", "label::English (en)"])
        choices.append(["yes_no", "yes", "Haan", "Yes"])

    settings = workbook.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["Training Form", "training_form", "2026-09-01"])

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def make_continuity_workbook(second=False, english_changed=False):
    workbook = Workbook()
    survey = workbook.active
    survey.title = "survey"
    survey.append(["type", "name", "label::English (en)", "label::Punjabi(pu)"])
    q1_english = "Name changed" if english_changed else "Name"
    q1_punjabi = "Naam changed" if english_changed else "Naam updated"
    survey_rows = [
        ["text", "q1", q1_english, q1_punjabi],
        ["text", "q2", "Age", "Updated age" if second else "Age translation"],
    ]
    if second:
        survey_rows = [survey_rows[1], survey_rows[0], ["text", "q3", "New", "Nava"]]
    for row in survey_rows:
        survey.append(row)

    choices = workbook.create_sheet("choices")
    choices.append(["list_name", "name", "label::Hindi (hi)", "label::English (en)"])
    choice_rows = [
        ["list_a", "shared", "Shared A", "Shared A"],
        ["list_b", "shared", "Updated list B" if second else "Shared B", "Shared B"],
    ]
    if second:
        choice_rows.reverse()
    for row in choice_rows:
        choices.append(row)

    settings = workbook.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["Continuity Form", "continuity_form", "2026-09-01"])

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


class XLSFormParserErrorTestCase(unittest.TestCase):
    def test_corrupted_workbook_is_rejected(self):
        with self.assertRaises(XLSFormValidationError):
            parse_xlsform(io.BytesIO(b"not a workbook"))


if __name__ == "__main__":
    unittest.main()
