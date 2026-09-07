"""Exercise upload -> review -> reordered upload -> exported cell continuity."""

import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app import create_app
from app.extensions import db
from app.models import ChoiceItem, Language, SurveyItem, TranslationReview, User, XLSForm
from app.services.review_workflow import get_or_create_form_assignment, save_question_with_choices
from app.services.xlsform_exporter import export_reviewed_xlsform


def version_workbook(names, choices):
    workbook = Workbook()
    survey = workbook.active
    survey.title = "survey"
    survey.append([
        "type", "name", "label::English (en)", "label::Hindi (hi)",
        "hint::English (en)", "hint::Hindi (hi)",
    ])
    for name in names:
        survey.append(["select_one yesno", name, f"English {name}", f"Hindi {name}",
                       f"Hint {name}", f"Hindi hint {name}"])
    sheet = workbook.create_sheet("choices")
    sheet.append(["list_name", "name", "label::English (en)", "label::Hindi (hi)"])
    for name in choices:
        sheet.append(["yesno", name, name.title(), f"Hindi {name}"])
    settings = workbook.create_sheet("settings")
    settings.append(["form_id", "form_title"])
    settings.append(["continuity", "Continuity"])
    result = io.BytesIO()
    workbook.save(result)
    workbook.close()
    return result.getvalue()


class VersionExportTestCase(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.app = create_app(type("TestConfig", (), {
            "TESTING": True, "SECRET_KEY": "test", "WTF_CSRF_ENABLED": False,
            "SQLALCHEMY_DATABASE_URI": "sqlite:///:memory:",
            "SQLALCHEMY_TRACK_MODIFICATIONS": False,
            "UPLOAD_FOLDER": str(Path(self.temp.name) / "uploads"),
            "EXPORT_FOLDER": str(Path(self.temp.name) / "exports"),
        }))
        self.context = self.app.app_context()
        self.context.push()
        db.create_all()
        self.admin = User(username="admin", email="admin@example.com", role=User.ROLE_ADMIN)
        self.reviewer = User(username="reviewer", email="reviewer@example.com", role=User.ROLE_REVIEWER)
        self.admin.set_password("test-password")
        self.reviewer.set_password("test-password")
        db.session.add_all([self.admin, self.reviewer])
        db.session.commit()
        self.client = self.app.test_client()
        with self.client.session_transaction() as session:
            session["_user_id"] = str(self.admin.id)
            session["_fresh"] = True

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.context.pop()
        self.temp.cleanup()

    def upload(self, content):
        response = self.client.post("/admin/xlsforms", data={
            "xlsform_file": (io.BytesIO(content), "version.xlsx"),
        }, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 200)
        return XLSForm.query.order_by(XLSForm.id.desc()).first()

    def test_inherited_edits_export_at_current_rows_and_keep_sources_intact(self):
        first_bytes = version_workbook(["removed", "kept", "shared"], ["yes", "no"])
        first = self.upload(first_bytes)
        language = Language.query.filter_by(language_code="hi").one()
        item = SurveyItem.query.filter_by(xlsform_id=first.id, name="kept").one()
        choice = ChoiceItem.query.filter_by(xlsform_id=first.id, name="yes").one()
        label = '<span style="color: red;"><b>Edited label</b></span>'
        hint = "<b>Edited hint</b>"
        answer = "<i>Edited yes</i>"
        save_question_with_choices(self.reviewer, first, language, item.id, label,
                                   {choice.id: answer}, {"hint": hint})
        db.session.commit()

        second_bytes = version_workbook(["kept", "added", "shared"], ["no", "yes", "maybe"])
        second = self.upload(second_bytes)
        assignment = get_or_create_form_assignment(self.reviewer, second, language)
        db.session.commit()
        exported = export_reviewed_xlsform(assignment, self.app.config["UPLOAD_FOLDER"],
                                          self.app.config["EXPORT_FOLDER"])
        workbook = load_workbook(exported)
        try:
            self.assertIn("Edited label", workbook["survey"]["D2"].value)
            self.assertIn("color:red", workbook["survey"]["D2"].value.replace(" ", ""))
            self.assertIn(hint, workbook["survey"]["F2"].value)
            self.assertEqual(workbook["survey"]["D3"].value, "Hindi added")
            self.assertEqual(workbook["choices"]["D2"].value, "Hindi no")
            self.assertIn(answer, workbook["choices"]["D3"].value)
            self.assertEqual(workbook["choices"]["D4"].value, "Hindi maybe")
            self.assertEqual(workbook["survey"].max_row, 4)
        finally:
            workbook.close()
        for form, original in ((first, first_bytes), (second, second_bytes)):
            self.assertEqual((Path(self.app.config["UPLOAD_FOLDER"]) / form.stored_filename).read_bytes(), original)

    def test_ambiguous_or_blank_identities_do_not_inherit(self):
        cases = [
            ("survey", ["kept", "kept"], ["kept"]),
            ("survey", ["kept"], ["kept", "kept"]),
            ("survey", [""], [""]),
            ("choices", ["yes", "yes"], ["yes"]),
            ("choices", ["yes"], ["yes", "yes"]),
            ("choices", [""], [""]),
        ]
        for sheet, before, after in cases:
            with self.subTest(sheet=sheet, before=before, after=after):
                source = self.upload(version_workbook(
                    before if sheet == "survey" else ["kept"],
                    before if sheet == "choices" else ["yes"],
                ))
                language = Language.query.filter_by(language_code="hi").one()
                for review in TranslationReview.query.filter_by(
                    xlsform_id=source.id, sheet_name=sheet, language_id=language.id,
                ).all():
                    review.edited_translation = "Must not carry ambiguous edit"
                    review.edited_by = self.reviewer.id
                db.session.commit()
                target = self.upload(version_workbook(
                    after if sheet == "survey" else ["kept"],
                    after if sheet == "choices" else ["yes"],
                ))
                reviews = TranslationReview.query.filter_by(
                    xlsform_id=target.id, sheet_name=sheet, language_id=language.id,
                ).all()
                self.assertTrue(reviews)
                self.assertTrue(all(review.edited_by is None for review in reviews))
                self.assertTrue(all(review.edited_translation != "Must not carry ambiguous edit" for review in reviews))

    def test_same_upload_reimport_uses_names_after_rows_move(self):
        source = self.upload(version_workbook(["removed", "kept"], ["yes", "no"]))
        language = Language.query.filter_by(language_code="hi").one()
        item = SurveyItem.query.filter_by(xlsform_id=source.id, name="kept").one()
        choice = ChoiceItem.query.filter_by(xlsform_id=source.id, name="yes").one()
        save_question_with_choices(self.reviewer, source, language, item.id,
                                   "Saved kept", {choice.id: "Saved yes"}, {"hint": "Saved hint"})
        db.session.commit()
        path = Path(self.app.config["UPLOAD_FOLDER"]) / source.stored_filename
        path.write_bytes(version_workbook(["kept", "added"], ["no", "yes"]))
        response = self.client.post(f"/admin/xlsforms/{source.id}/import")
        self.assertEqual(response.status_code, 302)
        for sheet, row, field, expected in [
            ("survey", 2, "label", "Saved kept"),
            ("survey", 2, "hint", "Saved hint"),
            ("choices", 3, "label", "Saved yes"),
        ]:
            review = TranslationReview.query.filter_by(
                xlsform_id=source.id, sheet_name=sheet, row_number=row,
                field_name=field, language_id=language.id,
            ).one()
            self.assertEqual(review.edited_translation, expected)
