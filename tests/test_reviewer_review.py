import io
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from app import create_app
from app.extensions import db
from app.models import (
    Language,
    ChoiceItem,
    ReviewerFormAssignment,
    SurveyItem,
    TranslationHistory,
    TranslationReview,
    User,
    UserLanguageAssignment,
    XLSForm,
)
from app.services.xlsform_importer import ensure_detected_languages, import_questionnaire_content
from app.services.display_text import clean_html_for_display, clean_text_for_display
from app.services.datetime_display import format_datetime_for_display
from app.services.review_workflow import display_question_label
from app.services.xlsform_parser import parse_xlsform


class TestConfig:
    TESTING = True
    SECRET_KEY = "test-secret"
    SQLALCHEMY_DATABASE_URI = "sqlite:///:memory:"
    SQLALCHEMY_TRACK_MODIFICATIONS = False
    WTF_CSRF_ENABLED = False
    UPLOAD_FOLDER = "instance/test_uploads"
    EXPORT_FOLDER = "instance/test_exports"


class CsrfTestConfig(TestConfig):
    WTF_CSRF_ENABLED = True


class CompactReviewerWorkflowTestCase(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.app = create_app(TestConfig)
        self.app.config["UPLOAD_FOLDER"] = str(Path(self.temp_dir.name) / "uploads")
        self.app.config["EXPORT_FOLDER"] = str(Path(self.temp_dir.name) / "exports")
        self.client = self.app.test_client()
        self.app_context = self.app.app_context()
        self.app_context.push()
        db.create_all()
        self.admin = self.create_user("admin", "admin@example.com", User.ROLE_ADMIN)
        self.reviewer = self.create_user("reviewer", "reviewer@example.com", User.ROLE_REVIEWER)
        self.other_reviewer = self.create_user("other", "other@example.com", User.ROLE_REVIEWER)
        self.english_reviewer = self.create_user("english", "english@example.com", User.ROLE_REVIEWER)
        self.xlsform, self.hindi, self.tamil, self.english = self.create_imported_xlsform()
        db.session.add(UserLanguageAssignment(user_id=self.reviewer.id, language_id=self.hindi.id))
        db.session.add(UserLanguageAssignment(user_id=self.other_reviewer.id, language_id=self.tamil.id))
        db.session.add(UserLanguageAssignment(user_id=self.english_reviewer.id, language_id=self.english.id))
        db.session.commit()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.app_context.pop()
        self.temp_dir.cleanup()

    def create_user(self, username, email, role, password="password123"):
        user = User(username=username, email=email, role=role, is_active=True)
        user.set_password(password)
        db.session.add(user)
        db.session.commit()
        return user

    def create_imported_xlsform(self):
        workbook_bytes = make_review_workbook()
        upload_dir = Path(self.app.config["UPLOAD_FOLDER"])
        upload_dir.mkdir(parents=True, exist_ok=True)
        stored_filename = "review.xlsx"
        (upload_dir / stored_filename).write_bytes(workbook_bytes.getvalue())

        parse_result = parse_xlsform(io.BytesIO(workbook_bytes.getvalue()))
        xlsform = XLSForm(
            original_filename="review.xlsx",
            stored_filename=stored_filename,
            uploaded_by=self.admin.id,
            form_title=parse_result.form_title,
            form_id=parse_result.form_id,
            version=parse_result.version,
            survey_row_count=parse_result.survey_row_count,
            choices_row_count=parse_result.choices_row_count,
            detected_language_count=len(parse_result.languages),
        )
        db.session.add(xlsform)
        ensure_detected_languages(parse_result.languages)
        db.session.flush()
        import_questionnaire_content(xlsform, parse_result)
        db.session.commit()
        return (
            xlsform,
            Language.query.filter_by(excel_header="label::Hindi (hi)").first(),
            Language.query.filter_by(excel_header="label::Tamil (ta)").first(),
            Language.query.filter_by(excel_header="label::English (en)").first(),
        )

    def login(self, username="reviewer"):
        return self.client.post(
            "/login",
            data={"username_or_email": username, "password": "password123"},
        )

    def test_items_before_id10010_are_visible(self):
        self.login()
        response = self.client.get("/reviewer/review")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Question 1 English", response.data)
        self.assertIn(b"Question 9 English", response.data)
        self.assertIn(b"Prashn 9", response.data)

    def test_reviewer_form_starts_with_first_reviewable_item_in_row_order(self):
        self.login()
        response = self.client.get("/reviewer/review")
        first_question = response.data.index(b"Question 1 English")

        self.assertLess(first_question, response.data.index(b"Question 2 English"))
        self.assertLess(response.data.index(b"Question 9 English"), response.data.index(b"Interview language"))

    def test_form_does_not_depend_on_id10010_name(self):
        SurveyItem.query.filter_by(name="Id10010").first().name = "MissingStart"
        db.session.commit()

        self.login()
        response = self.client.get("/reviewer/review")

        self.assertEqual(response.status_code, 200)
        self.assertNotIn(b"Review start question", response.data)
        self.assertIn(b"Question 1 English", response.data)
        self.assertIn(b"Interview language", response.data)

    def test_dashboard_count_includes_complete_form(self):
        self.login()
        response = self.client.get("/reviewer/")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Reviewable questions", response.data)
        self.assertIn(b'<dd class="col-sm-8">20</dd>', response.data)

    def test_question_name_is_not_visible_with_question_text(self):
        self.login()
        response = self.client.get("/reviewer/review")
        item = SurveyItem.query.filter_by(name="Id10016").first()
        row_html = self.question_row_html(response, item)

        self.assertNotIn(b"Variable", response.data)
        self.assertNotIn(b"compact-question-id", response.data)
        self.assertNotIn(b"Id10016", row_html)
        self.assertIn(b"Question 16 English", row_html)
        self.assertNotIn(b"Id10016 - Question 16 English", row_html)

    def test_repeated_headings_are_absent_and_main_values_are_read_only(self):
        self.login()
        response = self.client.get("/reviewer/review")

        self.assertNotIn(b"type=\"search\"", response.data)
        self.assertNotIn(b"Search</button>", response.data)
        self.assertNotIn(b"Excel row", response.data)
        self.assertNotIn(b"English Reference", response.data)
        self.assertNotIn(b"Hindi Translation", response.data)
        self.assertNotIn(b'name="question_16"', response.data)
        self.assertIn(b"Edit</button>", response.data)
        self.assertEqual(response.data.count(b'id="editTranslationModal"'), 1)

    def test_english_question_links_to_translation_for_assigned_language(self):
        self.login()
        response = self.client.get("/reviewer/review")

        self.assertIn(b'class="translate-action"', response.data)
        self.assertIn(b"<span>Translate</span>", response.data)
        self.assertIn(b"https://translate.google.com/?sl=en&amp;tl=hi&amp;text=Question%201%20English&amp;op=translate", response.data)
        self.assertIn(b'target="_blank"', response.data)

    def test_application_heading_displays_icmrva_tool_review(self):
        self.login()
        response = self.client.get("/reviewer/review")

        self.assertEqual(response.data.count(b"ICMRVA Tool Review"), 2)
        self.assertEqual(response.data.count(b'class="navbar-brand"'), 1)
        self.assertNotIn(b'<h1 class="h5 mb-0">ICMRVA Tool Review</h1>', response.data)
        self.assertNotIn(b"XLSForm Review</a>", response.data)
        self.assertNotIn(b"XLSForm Translation Review Portal", response.data)

    def test_top_and_bottom_finish_save_buttons_use_same_action(self):
        self.login()
        response = self.client.get("/reviewer/review")

        self.assertIn(b'id="finish-form-top" method="post"', response.data)
        self.assertIn(b'id="finish-form-bottom" method="post"', response.data)
        self.assertEqual(response.data.count(b"Finish / Save Form</button>"), 2)
        self.assertIn(b"Print Complete Form</button>", response.data)

    def test_choice_translations_are_displayed_compactly(self):
        self.login()
        response = self.client.get("/reviewer/review")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Male", response.data)
        self.assertIn(b"Purush", response.data)
        self.assertIn(b"Female", response.data)
        self.assertIn(b"Mahila", response.data)

    def test_select_one_choices_render_as_disabled_radio_buttons(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10017").first()
        response = self.client.get("/reviewer/review")
        row_html = self.question_row_html(response, item)

        self.assertEqual(response.status_code, 200)
        self.assertIn(f'id="question-row-{item.id}"'.encode(), response.data)
        self.assertEqual(row_html.count(b'type="radio" disabled'), 2)
        self.assertNotIn(b"English Option", response.data)
        self.assertNotIn(b"Hindi Option", response.data)

    def test_hindi_reviewer_sees_english_and_hindi_interview_languages_only(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10010").first()
        response = self.client.get("/reviewer/review")
        row_html = self.question_row_html(response, item)

        self.assertIn(b"Interview language", row_html)
        self.assertIn(b"English", row_html)
        self.assertIn(b"Hindi", row_html)
        self.assertNotIn(b"Tamil", row_html)
        self.assertNotIn(b"Marathi", row_html)
        self.assertNotIn(b"Punjabi", row_html)
        self.assertEqual(row_html.count(b'type="radio" disabled'), 2)

    def test_tamil_reviewer_gets_english_when_assigned_language_is_not_in_source_choices(self):
        self.login("other")
        item = SurveyItem.query.filter_by(name="Id10010").first()
        response = self.client.get("/reviewer/review")
        row_html = self.question_row_html(response, item)

        self.assertIn(b"Interview language", row_html)
        self.assertIn(b"English", row_html)
        self.assertNotIn(b"Tamil", row_html)
        self.assertNotIn(b"Hindi", row_html)
        self.assertNotIn(b"Marathi", row_html)
        self.assertNotIn(b"Punjabi", row_html)
        self.assertEqual(row_html.count(b'type="radio" disabled'), 1)

    def test_english_reviewer_sees_english_interview_language_once(self):
        self.login("english")
        item = SurveyItem.query.filter_by(name="Id10010").first()
        response = self.client.get("/reviewer/review")
        row_html = self.question_row_html(response, item)

        self.assertIn(b"Interview language", row_html)
        self.assertEqual(row_html.count(b'disabled>English (English)'), 1)
        self.assertNotIn(b"Hindi", row_html)
        self.assertNotIn(b"Tamil", row_html)
        self.assertEqual(row_html.count(b'type="radio" disabled'), 1)

    def test_interview_language_filtering_does_not_apply_to_other_select_one_questions(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10017").first()
        response = self.client.get("/reviewer/review")
        row_html = self.question_row_html(response, item)

        self.assertIn(b"Male", row_html)
        self.assertIn(b"Female", row_html)
        self.assertEqual(row_html.count(b'type="radio" disabled'), 2)

    def test_edit_modal_uses_same_filtered_interview_language_options(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10010").first()
        response = self.client.get("/reviewer/review")
        row_html = self.question_row_html(response, item)

        self.assertIn(b'data-choices=', row_html)
        self.assertIn(b'"english": "English (English)"', row_html)
        self.assertIn(b'"english": "Hindi (', row_html)
        self.assertNotIn(b'"english": "Tamil (', row_html)
        self.assertNotIn(b'"english": "Marathi (', row_html)

    def test_hidden_interview_language_options_remain_in_database(self):
        self.login()
        self.client.get("/reviewer/review")
        choices = ChoiceItem.query.filter_by(xlsform_id=self.xlsform.id, list_name="language").all()

        self.assertEqual({choice.name for choice in choices}, {"1", "2", "3", "4", "5", "6"})

    def test_hidden_interview_language_choice_submission_is_ignored(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10010").first()
        hidden_choice = ChoiceItem.query.filter_by(xlsform_id=self.xlsform.id, list_name="language", name="6").first()

        response = self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "Sakshatkar bhasha", f"choice_{hidden_choice.id}": "Hidden Malayalam edit"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertNotEqual(self.review("choices", hidden_choice.row_number).edited_translation, "Hidden Malayalam edit")

    def test_select_multiple_choices_render_as_disabled_checkboxes(self):
        self.login()
        response = self.client.get("/reviewer/review")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data.count(b'type="checkbox" disabled'), 3)
        self.assertIn(b"Fever", response.data)
        self.assertIn(b"Bukhar", response.data)
        self.assertIn(b"Cough", response.data)
        self.assertIn(b"Khansi", response.data)

    def test_answer_state_is_never_saved(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10017").first()
        response = self.client.get("/reviewer/review")

        self.assertNotIn(b'checked', response.data)
        self.assertNotIn(b'name="answer_', response.data)
        self.assertNotIn(b'type="radio" disabled name=', response.data)
        self.assertNotIn(b'type="checkbox" disabled name=', response.data)
        self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "Prashn 17", "answer_1": "male"},
        )

        self.assertEqual(self.review("choices", 2).edited_translation, "Purush")
        self.assertEqual(self.review("choices", 3).edited_translation, "Mahila")
        self.assertIsNone(self.review("choices", 2).edited_by)
        self.assertIsNone(self.review("choices", 3).edited_by)

    def test_review_uses_compact_bilingual_table(self):
        self.login()
        response = self.client.get("/reviewer/review")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'class="table table-sm table-bordered align-middle review-table mb-0"', response.data)
        self.assertIn(b"Options (English)", response.data)
        self.assertIn(b"Options (Hindi)", response.data)
        self.assertNotIn(b"Main Group English", response.data)
        self.assertNotIn(b"Pramukh Samuh", response.data)
        self.assertNotIn(b"Nested Group English", response.data)
        self.assertNotIn(b"Antarik Samuh", response.data)
        self.assertNotIn(b"begin_group", response.data)
        self.assertNotIn(b"end_group", response.data)
        self.assertNotIn(b"&lt;div&gt;Nested Group English&lt;/div&gt;", response.data)
        self.assertNotIn(b"&lt;p&gt;Nested Group English", response.data)

    def test_group_rows_do_not_interrupt_question_order(self):
        self.login()
        response = self.client.get("/reviewer/review")
        question_16 = response.data.index(b"Question 16 English")
        question_17 = response.data.index(b"Question 17 English")
        question_18 = response.data.index(b"Question 18 English")

        self.assertNotIn(b"Nested Group English", response.data)
        self.assertLess(question_17, question_18)
        self.assertLess(question_16, question_17)

    def test_modal_button_contains_clean_question_text_without_visible_name(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10016").first()
        response = self.client.get("/reviewer/review")

        self.assertIn(f'data-question-id="{item.id}"'.encode(), response.data)
        self.assertNotIn(b'data-question-name=', response.data)
        self.assertNotIn(b'id="modal-question-name"', response.data)
        self.assertNotIn(b"Question ID", response.data)
        self.assertIn(b"Question 16 English", response.data)
        self.assertNotIn(b"${Id10016}", response.data)
        self.assertNotIn(b"&lt;br&gt;", response.data)

    def test_modal_save_updates_question_translation(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10016").first()
        response = self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "Edited question 16"},
        )
        review = self.review("survey", item.row_number)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(review.edited_translation, "Edited question 16")
        self.assertTrue(response.get_json()["is_edited"])

    def test_question_edit_creates_history_entry(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10016").first()
        self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "Edited question 16"},
        )
        history = self.history_entries()

        self.assertEqual(len(history), 1)
        self.assertEqual(history[0].question_id, "Id10016")
        self.assertEqual(history[0].item_type, "Question")
        self.assertEqual(history[0].old_value, "Prashn 16")
        self.assertEqual(history[0].new_value, "Edited question 16")

    def test_modal_save_updates_choices(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10017").first()
        response = self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={
                "question_translation": "Edited select",
                "choice_1": "Edited male",
                "choice_2": "Edited female",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.review("choices", 2).edited_translation, "Edited male")
        self.assertEqual(self.review("choices", 3).edited_translation, "Edited female")
        self.assertTrue(response.get_json()["is_edited"])

    def test_choice_edit_creates_history_entry(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10017").first()
        response = self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "Prashn 17", "choice_1": "Edited male"},
        )
        history = self.history_entries()

        self.assertTrue(response.get_json()["is_edited"])
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0].question_id, "Id10017")
        self.assertEqual(history[0].item_type, "Choice")
        self.assertEqual(history[0].old_value, "Purush")
        self.assertEqual(history[0].new_value, "Edited male")

    def test_unchanged_save_does_not_create_fake_history(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10016").first()
        self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "Prashn 16"},
        )

        self.assertEqual(self.history_entries(), [])

    def test_multiple_edits_create_multiple_history_records_newest_first(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10016").first()
        self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "First edit"},
        )
        self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "Second edit"},
        )
        response = self.client.get("/reviewer/review")
        history = self.history_entries()

        self.assertEqual(len(history), 2)
        self.assertEqual(history[0].new_value, "Second edit")
        self.assertEqual(history[1].new_value, "First edit")
        self.assertLess(response.data.index(b"Second edit"), response.data.index(b"First edit"))

    def test_reviewer_sees_only_own_edit_history(self):
        self.login("other")
        item = SurveyItem.query.filter_by(name="Id10016").first()
        self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "Other reviewer edit"},
        )
        self.client.post("/logout")

        self.login()
        response = self.client.get("/reviewer/review")

        self.assertIn(b"Edit History (0 changes)", response.data)
        self.assertNotIn(b"Other reviewer edit", response.data)

    def test_history_is_filtered_by_xlsform_and_language(self):
        older_xlsform = XLSForm(
            original_filename="old.xlsx",
            stored_filename="old.xlsx",
            uploaded_by=self.admin.id,
            uploaded_at=datetime(2000, 1, 1, tzinfo=timezone.utc),
        )
        db.session.add(older_xlsform)
        db.session.flush()
        db.session.add(
            TranslationHistory(
                user_id=self.reviewer.id,
                xlsform_id=older_xlsform.id,
                language_id=self.hindi.id,
                sheet_name="survey",
                row_number=10,
                item_type="Question",
                question_id="OldQuestion",
                english_value="Old question",
                old_value="Old",
                new_value="Wrong form history",
            )
        )
        db.session.add(
            TranslationHistory(
                user_id=self.reviewer.id,
                xlsform_id=self.xlsform.id,
                language_id=self.tamil.id,
                sheet_name="survey",
                row_number=18,
                item_type="Question",
                question_id="Id10016",
                english_value="Tamil question",
                old_value="Old",
                new_value="Wrong language history",
            )
        )
        db.session.commit()

        self.login()
        response = self.client.get("/reviewer/review")

        self.assertIn(b"Edit History (0 changes)", response.data)
        self.assertNotIn(b"Wrong form history", response.data)
        self.assertNotIn(b"Wrong language history", response.data)

    def test_history_display_is_cleaned(self):
        db.session.add(
            TranslationHistory(
                user_id=self.reviewer.id,
                xlsform_id=self.xlsform.id,
                language_id=self.hindi.id,
                sheet_name="survey",
                row_number=18,
                item_type="Question",
                question_id="Id10016",
                english_value="<b>${Id10016} Id10016 - Question 16 English</b>",
                old_value="<span>${old_ref} Purana</span>",
                new_value="<span>${new_ref} Naya</span>",
            )
        )
        db.session.commit()

        self.login()
        response = self.client.get("/reviewer/review")

        self.assertIn(b"Question 16 English", response.data)
        self.assertIn(b"Purana", response.data)
        self.assertIn(b"Naya", response.data)
        self.assertNotIn(b"${old_ref}", response.data)
        self.assertNotIn(b"${new_ref}", response.data)

    def test_editing_question_translation_marks_question_edited(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10016").first()
        self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "Edited question 16"},
        )
        response = self.client.get("/reviewer/review")

        self.assertIn(f'id="question-row-{item.id}"'.encode(), response.data)
        self.assertIn(b"compact-question-edited", response.data)
        self.assertIn(f'id="edited-badge-{item.id}"'.encode(), response.data)

    def test_editing_choice_translation_marks_parent_question_edited_after_reload(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10017").first()
        self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "Prashn 17", "choice_1": "Edited male"},
        )
        response = self.client.get("/reviewer/review")

        self.assertIn(f'id="question-row-{item.id}"'.encode(), response.data)
        self.assertIn(f'id="edited-badge-{item.id}"'.encode(), response.data)
        self.assertIn(b"Edited male", response.data)

    def test_unchanged_questions_are_not_highlighted(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10016").first()
        response = self.client.get("/reviewer/review")

        row_html = self.question_row_html(response, item)
        self.assertNotIn(b"compact-question-edited", row_html)
        self.assertIn(f'id="edited-badge-{item.id}"'.encode(), response.data)
        self.assertIn(b"edited-badge d-none", response.data)

    def test_restoring_original_value_removes_edited_status(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10016").first()
        self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "Edited question 16"},
        )
        response = self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "Prashn 16"},
        )
        page = self.client.get("/reviewer/review")

        self.assertFalse(response.get_json()["is_edited"])
        row_html = self.question_row_html(page, item)
        self.assertNotIn(b"compact-question-edited", row_html)

    def test_reviewer_cannot_edit_another_language(self):
        self.login("other")
        item = SurveyItem.query.filter_by(name="Id10016").first()
        self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "Tamil edit"},
        )

        hindi_review = self.review("survey", item.row_number, self.hindi)
        tamil_review = self.review("survey", item.row_number, self.tamil)
        self.assertNotEqual(hindi_review.edited_translation, "Tamil edit")
        self.assertEqual(tamil_review.edited_translation, "Tamil edit")

    def test_first_question_can_be_edited_and_exported(self):
        self.login()
        first_item = SurveyItem.query.filter_by(name="Id10001").first()
        response = self.client.post(
            f"/reviewer/review/{first_item.id}/save",
            data={"question_translation": "First question edit"},
        )
        self.assertEqual(response.status_code, 200)

        self.client.post("/reviewer/review", data={})
        assignment = ReviewerFormAssignment.query.filter_by(user_id=self.reviewer.id).first()
        self.client.post("/logout")
        self.login("admin")
        response = self.client.get(f"/admin/assignments/{assignment.id}/download")
        exported = load_workbook(io.BytesIO(response.data), data_only=True)

        self.assertEqual(exported["survey"]["D3"].value, "Question 1 English\nFirst question edit")

    def test_preexisting_first_row_edit_is_exported(self):
        self.login()
        hidden_item = SurveyItem.query.filter_by(name="Id10001").first()
        hidden_review = self.review("survey", hidden_item.row_number)
        hidden_review.edited_translation = "Old hidden edit"
        hidden_review.edited_by = self.reviewer.id
        self.client.post("/reviewer/review", data={})
        assignment = ReviewerFormAssignment.query.filter_by(user_id=self.reviewer.id).first()
        db.session.commit()

        self.client.post("/logout")
        self.login("admin")
        response = self.client.get(f"/admin/assignments/{assignment.id}/download")
        exported = load_workbook(io.BytesIO(response.data), data_only=True)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(exported["survey"]["D3"].value, "Question 1 English\nOld hidden edit")

    def test_admin_export_contains_popup_saved_changes(self):
        self.login()
        item = SurveyItem.query.filter_by(name="Id10016").first()
        self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "Popup saved"},
        )
        assignment = ReviewerFormAssignment.query.filter_by(user_id=self.reviewer.id).first()
        original_path = Path(self.app.config["UPLOAD_FOLDER"]) / self.xlsform.stored_filename
        original_bytes = original_path.read_bytes()

        self.client.post("/logout")
        self.login("admin")
        response = self.client.get(f"/admin/assignments/{assignment.id}/download")
        exported = load_workbook(io.BytesIO(response.data), data_only=True)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(exported["survey"]["D18"].value, "${Id10016} Question 16 English\nPopup saved")
        self.assertEqual(exported["survey"]["E18"].value, "Question 16 English\nTamil 16")
        self.assertEqual(exported["choices"]["D2"].value, "Male\nPurush")
        self.assertEqual(exported["choices"]["D3"].value, "Female\nMahila")
        self.assertEqual(original_path.read_bytes(), original_bytes)

    def test_admin_download_returns_xlsx_attachment(self):
        self.login()
        self.client.post("/reviewer/review", data={})
        assignment = ReviewerFormAssignment.query.filter_by(user_id=self.reviewer.id).first()

        self.client.post("/logout")
        self.login("admin")
        response = self.client.get(f"/admin/assignments/{assignment.id}/download")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.content_type,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertTrue(response.headers["Content-Disposition"].lower().endswith(".xlsx"))

    def test_admin_download_missing_source_is_controlled(self):
        self.login()
        self.client.post("/reviewer/review", data={})
        assignment = ReviewerFormAssignment.query.filter_by(user_id=self.reviewer.id).first()
        (Path(self.app.config["UPLOAD_FOLDER"]) / self.xlsform.stored_filename).unlink()

        self.client.post("/logout")
        self.login("admin")
        response = self.client.get(f"/admin/assignments/{assignment.id}/download", follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Unable to generate the reviewed XLSForm", response.data)

    def test_edit_modal_scrolls_and_keeps_save_footer(self):
        self.login()
        response = self.client.get("/reviewer/review")
        css = (Path("app") / "static" / "css" / "app.css").read_text(encoding="utf-8")

        self.assertIn(b"modal-dialog modal-dialog-scrollable modal-lg review-edit-modal-dialog", response.data)
        self.assertIn(b">Save</button>", response.data)
        self.assertIn("max-height: 90vh", css)
        self.assertIn("max-height: calc(90vh - 140px)", css)
        self.assertIn("overflow-y: auto", css)
        self.assertIn("position: sticky", css)
        self.assertIn("word-break: break-word", css)

    def test_admin_download_missing_target_language_header_is_controlled(self):
        self.login()
        self.client.post("/reviewer/review", data={})
        assignment = ReviewerFormAssignment.query.filter_by(user_id=self.reviewer.id).first()
        source_path = Path(self.app.config["UPLOAD_FOLDER"]) / self.xlsform.stored_filename
        workbook = load_workbook(source_path)
        workbook["survey"]["D1"] = "label::Missing (xx)"
        workbook.save(source_path)

        self.client.post("/logout")
        self.login("admin")
        response = self.client.get(f"/admin/assignments/{assignment.id}/download", follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Unable to generate the reviewed XLSForm", response.data)

    def test_reviewer_display_cleans_html_from_survey_and_choices(self):
        self.login()
        response = self.client.get("/reviewer/review")

        self.assertIn(b"Question 16 English", response.data)
        self.assertNotIn(b"${Id10016}", response.data)
        self.assertIn(b"Prashn 16", response.data)
        self.assertNotIn(b"Nested Group English", response.data)
        self.assertNotIn(b"Antarik Samuh", response.data)
        self.assertIn(b"Fever", response.data)
        self.assertIn(b"Bukhar", response.data)
        self.assertIn(b"Pain", response.data)
        self.assertIn(b"Question 19 English", response.data)
        self.assertIn(b"Line 2", response.data)
        self.assertIn(b"Did have fever?", response.data)
        self.assertNotIn(b"${deceased_name}", response.data)
        self.assertNotIn(b"${age}", response.data)
        self.assertNotIn(b"${respondent_name}", response.data)
        self.assertNotIn(b"${group_ref}", response.data)
        self.assertNotIn(b"${symptom_ref}", response.data)
        self.assertNotIn(b"${name}", response.data)
        self.assertNotIn(b"&lt;b&gt;", response.data)
        self.assertNotIn(b"&lt;/b&gt;", response.data)
        self.assertNotIn(b"&lt;span", response.data)
        self.assertNotIn(b"&lt;/span&gt;", response.data)
        self.assertNotIn(b"&lt;br", response.data)
        self.assertNotIn(b"&lt;font", response.data)
        self.assertNotIn(b"alert(1)", response.data)
        self.assertNotIn(b"onclick", response.data)

    def test_original_database_and_workbook_values_keep_html_markup(self):
        original_path = Path(self.app.config["UPLOAD_FOLDER"]) / self.xlsform.stored_filename
        workbook_before = original_path.read_bytes()
        self.login()
        self.client.get("/reviewer/review")
        db.session.expire_all()

        item = SurveyItem.query.filter_by(name="Id10019").first()
        calculation_item = SurveyItem.query.filter_by(name="score").first()
        choice = ChoiceItem.query.filter_by(list_name="symptoms", name="fever").first()
        workbook = load_workbook(original_path, data_only=True)

        self.assertEqual(item.english_label, "Question 19 English<br>Line 2 ${age}")
        self.assertEqual(item.raw_row_data["label::Hindi (hi)"], "Question 19 English<br>Line 2 ${age}\n<p>Prashn 19 ${respondent_name}</p>")
        self.assertEqual(item.relevant, "${age} > 0")
        self.assertEqual(calculation_item.raw_row_data["calculation"], "${age} + 1")
        self.assertEqual(choice.english_label, "<font>Fever</font>")
        self.assertEqual(choice.raw_row_data["label::Hindi (hi)"], "<font>Fever</font>\n<span>Bukhar ${symptom_ref}</span>")
        self.assertEqual(workbook["survey"]["C23"].value, "Question 19 English<br>Line 2 ${age}")
        self.assertEqual(workbook["survey"]["F23"].value, "${age} > 0")
        self.assertEqual(workbook["survey"]["H26"].value, "${age} + 1")
        self.assertEqual(workbook["choices"]["C4"].value, "<font>Fever</font>")
        self.assertEqual(original_path.read_bytes(), workbook_before)

    def test_export_does_not_globally_strip_xlsform_references(self):
        self.login()
        self.client.post("/reviewer/review", data={})
        assignment = ReviewerFormAssignment.query.filter_by(user_id=self.reviewer.id).first()

        self.client.post("/logout")
        self.login("admin")
        response = self.client.get(f"/admin/assignments/{assignment.id}/download")
        exported = load_workbook(io.BytesIO(response.data), data_only=True)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(exported["survey"]["C23"].value, "Question 19 English<br>Line 2 ${age}")
        self.assertEqual(exported["survey"]["D23"].value, "Question 19 English<br>Line 2 ${age}\n<p>Prashn 19 ${respondent_name}</p>")
        self.assertEqual(exported["survey"]["F23"].value, "${age} > 0")
        self.assertEqual(exported["survey"]["H26"].value, "${age} + 1")
        self.assertEqual(exported["choices"]["D4"].value, "<font>Fever</font>\n<span>Bukhar ${symptom_ref}</span>")

    def test_admin_export_retains_complete_language_choice_list(self):
        self.login()
        self.client.post("/reviewer/review", data={})
        assignment = ReviewerFormAssignment.query.filter_by(user_id=self.reviewer.id).first()
        original_path = Path(self.app.config["UPLOAD_FOLDER"]) / self.xlsform.stored_filename
        original_bytes = original_path.read_bytes()

        self.client.post("/logout")
        self.login("admin")
        response = self.client.get(f"/admin/assignments/{assignment.id}/download")
        exported = load_workbook(io.BytesIO(response.data), data_only=True)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            [exported["choices"][f"B{row}"].value for row in range(7, 13)],
            ["1", "2", "3", "4", "5", "6"],
        )
        self.assertEqual(original_path.read_bytes(), original_bytes)

    def test_finish_save_sets_form_level_status(self):
        self.login()
        response = self.client.post("/reviewer/review", data={}, follow_redirects=True)
        assignment = ReviewerFormAssignment.query.filter_by(
            user_id=self.reviewer.id,
            xlsform_id=self.xlsform.id,
            language_id=self.hindi.id,
        ).first()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(assignment.status, ReviewerFormAssignment.STATUS_SAVED)
        self.assertIsNotNone(assignment.last_saved_at)

    def test_finish_save_message_counts_current_saved_translation_changes(self):
        self.login()
        question = SurveyItem.query.filter_by(name="Id10016").first()
        select_question = SurveyItem.query.filter_by(name="Id10017").first()
        self.client.post(
            f"/reviewer/review/{question.id}/save",
            data={"question_translation": "Edited question 16"},
        )
        self.client.post(
            f"/reviewer/review/{select_question.id}/save",
            data={"question_translation": "Prashn 17", "choice_1": "Edited male"},
        )

        response = self.client.post("/reviewer/review", data={}, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        self.assertIn(
            b"Form review saved successfully. 2 translation changes are currently saved.",
            response.data,
        )
        self.assertNotIn(b"Changed fields:", response.data)

    def review(self, sheet_name, row_number, language=None):
        language = language or self.hindi
        return TranslationReview.query.filter_by(
            xlsform_id=self.xlsform.id,
            sheet_name=sheet_name,
            row_number=row_number,
            language_id=language.id,
        ).first()

    def history_entries(self, user=None, language=None, xlsform=None):
        user = user or self.reviewer
        language = language or self.hindi
        xlsform = xlsform or self.xlsform
        return TranslationHistory.query.filter_by(
            user_id=user.id,
            xlsform_id=xlsform.id,
            language_id=language.id,
        ).order_by(TranslationHistory.created_at.desc(), TranslationHistory.id.desc()).all()

    def question_row_html(self, response, item):
        marker = f'id="question-row-{item.id}"'.encode()
        marker_index = response.data.index(marker)
        row_start = response.data.rfind(b"<tr", 0, marker_index)
        row_end = response.data.index(b"</tr>", marker_index)
        return response.data[row_start:row_end]


class DisplayTextTestCase(unittest.TestCase):
    def test_safe_html_and_markdown_are_rendered(self):
        rendered = str(clean_html_for_display('## <span style="color:red">**INTRODUCTION**</span>'))

        self.assertEqual(rendered, '<span style="color:red"><strong>INTRODUCTION</strong></span>')

    def test_unsafe_html_is_removed_from_rich_text(self):
        rendered = str(clean_html_for_display('<span onclick="alert(1)">Safe</span><script>alert(1)</script>'))

        self.assertEqual(rendered, "<span>Safe</span>")

    def test_bold_text_displays_as_text(self):
        self.assertEqual(clean_text_for_display("<b>Question text</b>"), "Question text")

    def test_span_tags_are_hidden(self):
        self.assertEqual(clean_text_for_display('<span style="color:red">Text</span>'), "Text")

    def test_br_creates_readable_separation(self):
        self.assertEqual(clean_text_for_display("Line 1<br>Line 2"), "Line 1\nLine 2")

    def test_html_entities_decode(self):
        self.assertEqual(clean_text_for_display("A&amp;B&nbsp;&lt;5&gt;"), "A&B <5>")

    def test_script_and_event_handler_markup_cannot_execute_or_display(self):
        self.assertEqual(
            clean_text_for_display('<span onclick="alert(1)">Safe</span><script>alert(1)</script>'),
            "Safe",
        )

    def test_xlsform_reference_token_is_removed_from_display_text(self):
        self.assertEqual(clean_text_for_display("Did ${deceased_name} have fever?"), "Did have fever?")

    def test_multiple_xlsform_reference_tokens_are_removed_from_display_text(self):
        self.assertEqual(clean_text_for_display("${name} used ${age} years"), "used years")

    def test_normal_curly_braces_remain_in_display_text(self):
        self.assertEqual(clean_text_for_display("Price is {100} but ${price} hidden"), "Price is {100} but hidden")

    def test_question_id_prefix_is_removed_from_display_label_only(self):
        item = SimpleNamespace(name="Id10016", english_label="<b>Id10016</b> - What was the age?")

        self.assertEqual(display_question_label(item), "What was the age?")

    def test_question_id_space_prefix_is_removed_from_display_label_only(self):
        item = SimpleNamespace(name="Id10017", english_label="Id10017 Sex of deceased?")

        self.assertEqual(display_question_label(item), "Sex of deceased?")

    def test_ordinary_question_name_is_preserved_when_label_starts_with_it(self):
        item = SimpleNamespace(name="Site", english_label="Site of VA Interviewer")

        self.assertEqual(display_question_label(item), "Site of VA Interviewer")

    def test_utc_datetime_is_displayed_in_india_time(self):
        saved_at = datetime(2026, 9, 4, 16, 0, tzinfo=timezone.utc)

        self.assertEqual(
            format_datetime_for_display(saved_at, "Asia/Kolkata"),
            "04 Sep 2026, 09:30 PM IST",
        )


class CompactReviewerWorkflowCsrfTestCase(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.app = create_app(CsrfTestConfig)
        self.app.config["UPLOAD_FOLDER"] = str(Path(self.temp_dir.name) / "uploads")
        self.app.config["EXPORT_FOLDER"] = str(Path(self.temp_dir.name) / "exports")
        self.client = self.app.test_client()
        self.app_context = self.app.app_context()
        self.app_context.push()
        db.create_all()
        self.reviewer = User(username="reviewer", email="reviewer@example.com", role=User.ROLE_REVIEWER)
        self.reviewer.set_password("password123")
        db.session.add(self.reviewer)
        db.session.commit()
        parse_result = parse_xlsform(make_review_workbook())
        xlsform = XLSForm(
            original_filename="review.xlsx",
            stored_filename="review.xlsx",
            uploaded_by=self.reviewer.id,
            form_title=parse_result.form_title,
            form_id=parse_result.form_id,
            version=parse_result.version,
            survey_row_count=parse_result.survey_row_count,
            choices_row_count=parse_result.choices_row_count,
            detected_language_count=len(parse_result.languages),
        )
        db.session.add(xlsform)
        ensure_detected_languages(parse_result.languages)
        db.session.flush()
        import_questionnaire_content(xlsform, parse_result)
        hindi = Language.query.filter_by(excel_header="label::Hindi (hi)").first()
        db.session.add(UserLanguageAssignment(user_id=self.reviewer.id, language_id=hindi.id))
        db.session.commit()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.app_context.pop()
        self.temp_dir.cleanup()

    def test_csrf_protected_modal_save(self):
        item = SurveyItem.query.filter_by(name="Id10016").first()
        with self.client.session_transaction() as session:
            session["_user_id"] = str(self.reviewer.id)
            session["_fresh"] = True
        response = self.client.post(
            f"/reviewer/review/{item.id}/save",
            data={"question_translation": "Blocked"},
        )

        self.assertEqual(response.status_code, 400)

def make_review_workbook():
    workbook = Workbook()
    survey = workbook.active
    survey.title = "survey"
    survey.append(
        [
            "type",
            "name",
            "label::English (en)",
            "label::Hindi (hi)",
            "label::Tamil (ta)",
            "relevant",
            "appearance",
            "calculation",
        ]
    )
    survey.append(["begin_group", "main_group", "<b>Main Group English ${group_ref}</b>", "<span>Main Group English ${group_ref}\nPramukh Samuh ${group_ref}</span>", "Main Group English\nTamil Group", None, None])
    for index in range(1, 17):
        question_type = "select_one language" if index == 10 else "text"
        english_label = "Interview language" if index == 10 else f"Question {index} English"
        hindi_label = "Interview language\nSakshatkar bhasha" if index == 10 else f"Question {index} English\nPrashn {index}"
        tamil_label = "Interview language\nNerkanal mozhi" if index == 10 else f"Question {index} English\nTamil {index}"
        survey.append(
            [
                question_type,
                f"Id100{index:02d}",
                "${Id10016} Question 16 English" if index == 16 else english_label,
                "${Id10016} Question 16 English\nPrashn 16" if index == 16 else hindi_label,
                tamil_label,
                None,
                None,
            ]
        )
    survey.append(["begin_group", "nested_group", "<div>Nested Group English</div>", "<p>Nested Group English<br>Antarik Samuh</p>", "Nested Group English\nNested Tamil", None, None])
    survey.append(
        [
            "select_one sex",
            "Id10017",
            "Question 17 English",
            "Question 17 English\nPrashn 17",
            "Question 17 English\nTamil 17",
            None,
            None,
        ]
    )
    survey.append(["end_group", None, None, None, None, None, None])
    survey.append(
        [
            "text",
            "Id10018",
            "Question 18 English",
            "Question 18 English\nPrashn 18",
            "Question 18 English\nTamil 18",
            None,
            None,
        ]
    )
    survey.append(
        [
            "select_multiple symptoms",
            "Id10019",
            "Question 19 English<br>Line 2 ${age}",
            "Question 19 English<br>Line 2 ${age}\n<p>Prashn 19 ${respondent_name}</p>",
            "Question 19 English\nTamil 19",
            "${age} > 0",
            None,
        ]
    )
    survey.append(
        [
            "note",
            "Id10020",
            "Did ${deceased_name} have fever?",
            "Did ${deceased_name} have fever?\nBukhar tha?",
            "Did ${deceased_name} have fever?\nTamil fever",
            "${deceased_name} != ''",
            None,
        ]
    )
    survey.append(["end_group", None, None, None, None, None, None])
    survey.append(["calculate", "score", None, None, None, None, None, "${age} + 1"])

    choices = workbook.create_sheet("choices")
    choices.append(["list_name", "name", "label::English (en)", "label::Hindi (hi)", "label::Tamil (ta)"])
    choices.append(["sex", "male", "Male", "Male\nPurush", "Male\nAan"])
    choices.append(["sex", "female", "Female", "Female\nMahila", "Female\nPen"])
    choices.append(["symptoms", "fever", "<font>Fever</font>", "<font>Fever</font>\n<span>Bukhar ${symptom_ref}</span>", "Fever\nKaaichal"])
    choices.append(["symptoms", "cough", '<span onclick="alert(1)">Cough</span><script>alert(1)</script>', "Cough&nbsp;&amp;\nKhansi", "Cough\nIrumal"])
    choices.append(["symptoms", "pain", "${name} Pain ${age}", "${name} Pain ${age}\nDard", "Pain\nVali"])
    choices.append(["language", "1", "English (English)", "English\nAngrezi", "English"])
    choices.append(["language", "2", "Hindi (हिन्दी)", "Hindi\nHindi", "Hindi"])
    choices.append(["language", "3", "Marathi (मराठी)", "Marathi\nMarathi", "Marathi"])
    choices.append(["language", "4", "Bangla (বাংলা)", "Bangla\nBangla", "Bangla"])
    choices.append(["language", "5", "Kannada (ಕನ್ನಡ)", "Kannada\nKannada", "Kannada"])
    choices.append(["language", "6", "Malayalam (മലയാളം)", "Malayalam\nMalayalam", "Malayalam"])

    settings = workbook.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["Review Form", "review_form", "v1"])

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


if __name__ == "__main__":
    unittest.main()
